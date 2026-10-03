from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .dedupe import Listing

HEADERS = ["New", "First Seen", "Posted", "Company", "Role", "Location", "Country", "Remote",
           "Work Auth", "Sources", "Link"]
WIDTHS = [6, 12, 12, 28, 55, 45, 9, 8, 42, 30, 10]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
NEW_FILL = PatternFill("solid", fgColor="FFF2CC")


def write_excel(path, sheets: dict[str, list[tuple[Listing, str]]]):
    """sheets: {sheet title: [(listing, work-auth note), ...]}. New listings
    are listed first and highlighted on every sheet."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        _fill_sheet(wb.create_sheet(title), rows)
    wb.save(path)


def _fill_sheet(ws, rows):
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL

    rows = sorted(rows, key=lambda r: not r[0].is_new)  # stable: keeps newest-posted order within each group
    for item, note in rows:
        job = item.job
        ws.append([
            "NEW" if item.is_new else "", item.first_seen[:10], job.posted_on, job.company, job.title,
            job.location, (job.country or "").replace(",", ", "), "Yes" if job.remote else "",
            note, ", ".join(item.sources), "Open",
        ])
        if item.is_new:
            for cell in ws[ws.max_row]:
                cell.fill = NEW_FILL
        link = ws.cell(row=ws.max_row, column=len(HEADERS))
        link.hyperlink = job.url
        link.font = Font(color="0563C1", underline="single")

    for i, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
