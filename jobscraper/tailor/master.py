"""The master resume: every fact you might put on a resume, in one YAML file.

It can hold more bullets and projects than fit on a page; tailoring picks
and rewords from it. Nothing outside this file ever appears on a tailored
resume (see guard.py). Format: see master_resume.example.yaml.
"""

from pathlib import Path

import yaml

SECTIONS_WITH_BULLETS = ("experience", "projects", "leadership")


def load(path: Path) -> dict:
    data = yaml.safe_load(path.read_text())
    if not isinstance(data, dict) or not data.get("name"):
        raise ValueError(f"{path} doesn't look like a master resume (needs at least `name:`)")
    _assign_ids(data)
    return data


def save(data: dict, path: Path):
    path.write_text(yaml.safe_dump(data, sort_keys=False, allow_unicode=True, width=100))


def _assign_ids(data: dict):
    """Give every entry and bullet a stable id ("experience.1", "experience.1.b2")
    so tailored output can point back at the exact original it came from."""
    for section in SECTIONS_WITH_BULLETS:
        for i, entry in enumerate(data.get(section) or [], start=1):
            entry.setdefault("id", f"{section}.{i}")
            bullets = entry.get("bullets") or []
            entry["bullets"] = [
                b if isinstance(b, dict) else {"text": b} for b in bullets
            ]
            for j, bullet in enumerate(entry["bullets"], start=1):
                bullet.setdefault("id", f"{entry['id']}.b{j}")


def bullet_index(data: dict) -> dict[str, dict]:
    """id -> bullet for every bullet in the master."""
    return {b["id"]: b for s in SECTIONS_WITH_BULLETS for e in data.get(s) or [] for b in e["bullets"]}


def all_text(data: dict) -> str:
    """Everything in the master as one string, for the no-new-facts check."""
    return yaml.safe_dump(data, allow_unicode=True)


def read_resume_file(path: Path) -> str:
    """Plain text of an existing resume (PDF, DOCX, TXT/MD) for importing."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        import docx
        doc = docx.Document(str(path))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)
    if suffix in (".txt", ".md", ".tex"):
        return path.read_text()
    raise ValueError(f"can't read {suffix} files; export the resume as PDF, DOCX or TXT")
